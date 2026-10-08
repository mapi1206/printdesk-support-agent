"""Builds data.json (used by the app), CSV tables and an Excel workbook.

Run:  python3 data/generate.py
"""
import csv, json, random, datetime as dt, pathlib
from catalog import *

OUT = pathlib.Path(__file__).parent
rng = random.Random(42)
TODAY = dt.date(2026, 10, 5)
N_CASES = 480


def wchoice(pairs):
    items, weights = zip(*pairs)
    return rng.choices(items, weights=weights)[0]


def split_cat(pc):
    if not pc: return None, None
    c, _, h = pc.partition("|"); return c, (h or None)


def pick_part(pc, printer):
    cat, hint = split_cat(pc)
    cands = [p for p in PARTS if p["cat"] == cat and printer in p["compat"]]
    if hint and any(hint in p["name"].lower() for p in cands):
        cands = [p for p in cands if hint in p["name"].lower()]
    if not cands:
        return None
    off = [p for p in cands if p["official"]]
    alt = [p for p in cands if not p["official"]]
    pool = off if (off and (not alt or rng.random() < 0.68)) else alt
    pool.sort(key=lambda p: p["price"])
    return rng.choice(pool[:2])["id"]


cases = []
for i in range(N_CASES):
    printer = wchoice(PRINTER_POP.items())
    probs = [(p, p["freq"]) for p in PROBLEMS if printer in p["printers"]]
    prob = wchoice(probs)
    channel = wchoice(CHANNELS)
    lang = wchoice(LANGS)
    # cases get more frequent towards today (growing shop)
    age = int(abs(rng.triangular(0, 365, 0)))
    date = TODAY - dt.timedelta(days=age)
    resolved = rng.random() < 0.87
    steps = prob["steps"]
    solved_step = None
    custom = None
    if resolved:
        if rng.random() < 0.04:
            fixes = [f for f in CUSTOM_FIXES if f[0] == prob["id"]]
            if fixes:
                custom = rng.choice(fixes)[1]
        if custom is None:
            solved_step = rng.choices(range(len(steps)), weights=[s[0] for s in steps])[0]
    tried = (solved_step + 1) if solved_step is not None else (rng.randint(2, len(steps)) if not resolved else len(steps))
    part_id = None
    if solved_step is not None and steps[solved_step][3]:
        part_id = pick_part(steps[solved_step][3], printer)
    escalated = (not resolved) or (tried >= 4 and rng.random() < 0.3)
    if resolved:
        cr = 5 if tried <= 2 else (4 if tried <= 3 else rng.choice([3, 4]))
        if rng.random() < 0.15: cr = max(1, cr - 1)
    else:
        cr = rng.choice([1, 2, 2, 3])
    customer_rating = cr if rng.random() < 0.62 else None
    support_rating = (rng.choice([3, 4, 4, 5, 5]) if resolved else rng.choice([2, 3])) if channel != "chat" or escalated else None
    purchase_age = rng.randint(1, 40)
    sentiment = "frustrated" if (not resolved or tried >= 3) and rng.random() < 0.6 else wchoice([("neutral", 6), ("positive", 2), ("frustrated", 2), ("angry", 0.6)])
    urgency = "high" if sentiment in ("angry",) or (prob["cat"] in ("electronics",) and rng.random() < .4) else wchoice([("low", 4), ("medium", 5), ("high", 1)])
    cases.append(dict(
        id=f"H{i+1:04d}", source="history", date=date.isoformat(), printer=printer, problem=prob["id"],
        channel=channel, lang=lang, steps_tried=tried,
        solved_step=solved_step, custom_fix=custom, resolved=resolved, escalated=escalated,
        part_id=part_id, customer_rating=customer_rating, support_rating=support_rating,
        minutes=rng.randint(3, 12) * tried + (rng.randint(10, 60) if escalated else 0),
        purchase_months=purchase_age, warranty="in" if purchase_age <= 24 else "out",
        sentiment=sentiment, urgency=urgency,
    ))
cases.sort(key=lambda c: c["date"])

problems_out = [dict(id=p["id"], cat=p["cat"], printers=p["printers"], en=p["en"], de=p["de"], kw=p["kw"],
                     steps=[dict(id=f"{p['id']}-S{j+1}", en=s[1], de=s[2], part_cat=split_cat(s[3])[0], part_hint=split_cat(s[3])[1]) for j, s in enumerate(p["steps"])])
                for p in PROBLEMS]

# ---- simulated shop orders (order lookup for the agent); includes the order numbers used in the demo emails
FIRST=["Anna","Bence","Clara","David","Emma","Felix","Greta","Hannah","Ivan","Julia","Kristóf","Lena","Marco","Nora","Oskar","Petra","Quentin","Rita","Stefan","Tomasz","Ursula","Viktor","Wiebke","Zoltán"]
LAST=["Bauer","Kovács","Novak","Rossi","Dubois","Schmidt","Nagy","Horváth","Weber","Müller","Kowalski","Fischer","Moreau","Varga","Huber","Lehner"]
orders=[]
def mk_order(no, name, email, date, printer, combo, extra=()):
    pr=next(p for p in PRINTERS if p["id"]==printer)
    items=[dict(type="printer",id=printer,name=f"Bambu Lab {pr['name']}{' Combo' if combo else ''}",price=(pr['combo'] if combo else pr['price']) or 0)]
    for pid in extra:
        pt=next(p for p in PARTS if p["id"]==pid); items.append(dict(type="part",id=pid,name=pt["name"],price=pt["price"]))
    d=dt.date.fromisoformat(date)
    orders.append(dict(order_no=no,name=name,email=email,date=date,delivered=(d+dt.timedelta(days=3)).isoformat(),status="delivered",items=items,
        warranty_until=(d.replace(year=d.year+2)).isoformat()))
mk_order("3DJ-2026-048213","Szabó Péter","szabo.peter@example.com","2026-04-08","p1s",True,["BL-PLT-ENG"])
mk_order("3DJ-2024-077315","James Carter","james.carter@example.com","2024-03-12","x1c",True)
mk_order("3DJ-2219-8841","Markus Weber","markus.weber@example.com","2025-02-14","p1s",False)
mk_order("3DJ-2026-061190","Camille Laurent","camille.laurent@example.com","2026-06-17","p2s",True)
mk_order("3DJ-2025-151877","Tóth Eszter","toth.eszter@example.com","2025-12-03","a1mini",False,["BL-PWIPE-A1"])
mk_order("3DJ-2026-031544","Kertész Alma","kertesz.alma@example.com","2026-03-10","p1s",True)
mk_order("3DJ-2025-102233","Piotr Nowak","piotr.nowak@example.com","2025-08-21","p1s",False,["BL-PLT-TPEI"])
for i in range(70):
    printer=wchoice([(k,v) for k,v in PRINTER_POP.items() if next(p for p in PRINTERS if p["id"]==k)["price"]])
    nm=f"{rng.choice(FIRST)} {rng.choice(LAST)}"
    d=TODAY-dt.timedelta(days=rng.randint(5,900))
    extra=rng.sample([p["id"] for p in PARTS if printer in p["compat"]],k=rng.choice([0,0,1,2]))
    mk_order(f"3DJ-{d.year}-{rng.randint(100000,199999)}",nm,nm.lower().replace(" ",".").replace("ó","o").replace("ü","ue").replace("á","a")+"@example.com",d.isoformat(),printer,rng.random()<.55,extra)

data = dict(
    meta=dict(snapshot=SNAPSHOT_DATE, source=SOURCE, currency="EUR", free_shipping_from=FREE_SHIPPING_FROM,
              std_shipping=STD_SHIPPING, warranty_months=24,
              note="Prices, availability and compatibility from 3djake.at on 2026-10-05. Stock levels and case history are simulated."),
    printers=PRINTERS, parts=PARTS, categories={k: {"en": v[0], "de": v[1]} for k, v in CATEGORIES.items()},
    problems=problems_out, cases=cases, orders=orders,
)
(OUT / "data.json").write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")))

# ---- CSV tables
tables = OUT / "tables"; tables.mkdir(exist_ok=True)
def write_csv(name, rows, cols):
    with open(tables / name, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(cols)
        for r in rows: w.writerow([("|".join(r[c]) if isinstance(r.get(c), list) else r.get(c)) for c in cols])
    return rows, cols

sheets = {
    "printers": write_csv("printers.csv", PRINTERS, ["id", "name", "series", "enclosed", "plate", "hotend", "price", "combo", "released", "ams"]),
    "parts": write_csv("parts.csv", PARTS, ["id", "name", "brand", "official", "cat", "price", "was", "days", "stock", "reorder_level", "compat"]),
}
steps_rows = [dict(problem=p["id"], step=s["id"], order=j + 1, en=s["en"], de=s["de"], part_cat=s["part_cat"], part_hint=s["part_hint"]) for p in problems_out for j, s in enumerate(p["steps"])]
sheets["problems"] = write_csv("problems.csv", problems_out, ["id", "cat", "en", "de", "kw", "printers"])
sheets["solution_steps"] = write_csv("solution_steps.csv", steps_rows, ["problem", "step", "order", "en", "de", "part_cat", "part_hint"])
sheets["orders"] = write_csv("orders.csv", [dict(o,items=[i["name"] for i in o["items"]]) for o in orders], ["order_no","name","email","date","delivered","status","warranty_until","items"])
sheets["cases"] = write_csv("cases.csv", cases, list(cases[0].keys()))

# ---- Excel workbook (same tables)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
wb = Workbook(); wb.remove(wb.active)
for name, (rows, cols) in sheets.items():
    ws = wb.create_sheet(name)
    ws.append(cols)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F4858")
    for r in rows: ws.append([("|".join(r[c]) if isinstance(r.get(c), list) else r.get(c)) for c in cols])
    ws.freeze_panes = "A2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or "")) for c in col[:60]) + 2))
wb.save(OUT / "printdesk_data.xlsx")

res = sum(c["resolved"] for c in cases)
print(f"parts={len(PARTS)} problems={len(PROBLEMS)} cases={len(cases)} resolved={res/len(cases):.0%} size={(OUT/'data.json').stat().st_size//1024}KB")
